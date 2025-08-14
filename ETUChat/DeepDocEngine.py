import os
import json
import pandas as pd
import yaml
from bs4 import BeautifulSoup
from markdown import markdown
from PIL import Image

from deepdoc_pdfparser import parse_pdf
from docx import Document
from pptx import Presentation
import openpyxl


class DeepDocEngine:
    def __init__(self, ocr_lang='eng'):
        self.ocr_lang = ocr_lang

    # PDF
    def parse_pdf(self, file_path):
        doc = parse_pdf(file_path)
        return {"type": "pdf", "content": doc}

    # Word
    def parse_docx(self, file_path):
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        return {"type": "docx", "paragraphs": paragraphs}

    # PPT
    def parse_pptx(self, file_path):
        prs = Presentation(file_path)
        slides = []
        for slide_index, slide in enumerate(prs.slides, 1):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text")]
            slides.append({"slide_number": slide_index, "texts": texts})
        return {"type": "pptx", "slides": slides}

    # Excel
    def parse_xlsx(self, file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            sheets[sheet_name] = rows
        return {"type": "xlsx", "sheets": sheets}

    # TXT / TEXT
    def parse_txt(self, file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
        return {"type": "txt", "content": content}

    # JSON
    def parse_json(self, file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {"type": "json", "data": data}

    # CSV
    def parse_csv(self, file_path):
        df = pd.read_csv(file_path)
        return {"type": "csv", "data": df.values.tolist(), "columns": df.columns.tolist()}

    # YAML
    def parse_yaml(self, file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            data = yaml.safe_load(f)
        return {"type": "yaml", "data": data}

    # XML / HTML
    def parse_html_xml(self, file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
        soup = BeautifulSoup(content, 'lxml')
        return {"type": "html/xml", "text": soup.get_text()}

    # Markdown
    def parse_md(self, file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
        html = markdown(content)
        return {"type": "markdown", "html": html, "text": BeautifulSoup(html, 'lxml').get_text()}


    # 自动识别
    def parse(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return self.parse_pdf(file_path)
        elif ext == ".docx":
            return self.parse_docx(file_path)
        elif ext == ".pptx":
            return self.parse_pptx(file_path)
        elif ext == ".xlsx":
            return self.parse_xlsx(file_path)
        elif ext in [".txt", ".text"]:
            return self.parse_txt(file_path)
        elif ext == ".json":
            return self.parse_json(file_path)
        elif ext == ".csv":
            return self.parse_csv(file_path)
        elif ext in [".yml", ".yaml"]:
            return self.parse_yaml(file_path)
        elif ext in [".html", ".htm", ".xml"]:
            return self.parse_html_xml(file_path)
        elif ext == ".md":
            return self.parse_md(file_path)
        else:
            raise ValueError(f"不支持的文件类型: {ext}")
